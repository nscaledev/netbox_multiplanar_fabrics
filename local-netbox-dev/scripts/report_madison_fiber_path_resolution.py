from __future__ import annotations

import csv
import re
from collections import Counter, defaultdict
from pathlib import Path
import sys

from dcim.models import Device, FrontPort, Interface, Rack
from netbox_plant_graph.models import Endpoint


for candidate in (
    Path('/opt/netbox/local-plugins/netbox_multiplanar_fabrics/local-netbox-dev/scripts'),
    Path('/Users/mencken/github-repos/netbox_multiplanar_fabrics/local-netbox-dev/scripts'),
    Path(globals().get('__file__', '.')).resolve().parent,
):
    if candidate.exists() and str(candidate) not in sys.path:
        sys.path.insert(0, str(candidate))

from madison_v2_graph import (  # noqa: E402
    endpoint_for_device_port,
    get_fabric,
    mpo_endpoint_for_device_port,
)


MAD_SITE_SLUG = 'gs001'
FABRIC_NAME = 'GS001 RoCE Fabric'
SHUFFLE_CONTAINMENT_MARKER = 'madison_shuffle_flattened_containment'
SOURCE_WORKBOOK_PATHS = [
    Path('/opt/netbox/local-plugins/netbox_multiplanar_fabrics/local-netbox-dev/data/source/mencken-nc-backend-fiber-worksheet.xlsx'),
    Path('/Users/mencken/Documents/mencken-nc-backend-fiber-worksheet.xlsx'),
]
SOURCE_CSV_PATHS = [
    Path('/opt/netbox/local-plugins/netbox_multiplanar_fabrics/local-netbox-dev/data/generated/madison_fiber_worksheet_paths.csv'),
    Path('/Users/mencken/github-repos/netbox_multiplanar_fabrics/local-netbox-dev/data/generated/madison_fiber_worksheet_paths.csv'),
]
OUTPUT_PATHS = [
    Path('/opt/netbox/local-plugins/netbox_multiplanar_fabrics/local-netbox-dev/data/generated/madison_fiber_path_resolution.csv'),
    Path('/Users/mencken/github-repos/netbox_multiplanar_fabrics/local-netbox-dev/data/generated/madison_fiber_path_resolution.csv'),
]
BE_LEAF_PATTERN = re.compile(r'\bBE LEAF#(?P<switch>\d+)\.NIC(?P<nic>\d+)(?P<side>[AB])\.PL(?P<plane>\d+)\b')


def source_workbook_path() -> Path:
    for path in SOURCE_WORKBOOK_PATHS:
        if path.exists():
            return path
    raise RuntimeError(f'Corrected fiber worksheet not found in: {SOURCE_WORKBOOK_PATHS}')


def source_csv_path() -> Path | None:
    for path in SOURCE_CSV_PATHS:
        if path.exists():
            return path
    return None


def output_path() -> Path:
    for path in OUTPUT_PATHS:
        if path.parent.exists():
            return path
    OUTPUT_PATHS[-1].parent.mkdir(parents=True, exist_ok=True)
    return OUTPUT_PATHS[-1]


def read_worksheet_rows() -> list[dict[str, int]]:
    csv_path = source_csv_path()
    if csv_path is not None:
        with csv_path.open(newline='') as handle:
            return [
                {key: int(value) for key, value in row.items()}
                for row in csv.DictReader(handle)
            ]

    from openpyxl import load_workbook

    workbook = load_workbook(source_workbook_path(), read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = []
    found_header = False
    for values in worksheet.iter_rows(values_only=True):
        if values and values[0] == 'Per-role-MPO-index':
            found_header = True
            continue
        if not found_header or not values or values[0] is None:
            continue
        rows.append(
            {
                'path_index': int(values[0]),
                'pattern_su': int(values[1]),
                'nvl72_rack_ordinal': int(values[2]),
                'compute_tray_ordinal': int(values[3]),
                'cx8': int(values[4]),
                'gb300_mpo': int(values[5]),
                'shuffle_box_ordinal': int(values[7]),
                'shuffle_tray': int(values[8]),
                'shuffle_cassette': int(values[9]),
                'shuffle_mpo': int(values[10]),
                'plane': int(values[12]),
                'leaf_switch': int(values[13]),
                'leaf_cage': int(values[14]),
                'leaf_mpo': int(values[15]),
            }
        )
    return rows


def natural_rack_key(rack: Rack):
    match = re.search(r'^(?P<prefix>.*?)-C(?P<col>\d+)$', rack.name)
    if match:
        return (match.group('prefix'), int(match.group('col')), rack.name)
    return (rack.name, 0, rack.name)


def su_slug(tags) -> int | None:
    su_tags = sorted(tag.slug for tag in tags if tag.slug.startswith('nv_su_'))
    if len(su_tags) != 1:
        return None
    suffix = su_tags[0].rsplit('_', 1)[-1]
    return int(suffix) if suffix.isdigit() else None


def racks_by_su() -> dict[int, list[Rack]]:
    grouped = defaultdict(list)
    for rack in Rack.objects.filter(site__slug=MAD_SITE_SLUG).prefetch_related('tags').select_related('role'):
        su = su_slug(rack.tags.all())
        if su is not None:
            grouped[su].append(rack)
    return grouped


def gb300_racks_for_su(racks: list[Rack]) -> list[Rack]:
    return sorted([rack for rack in racks if rack.role and rack.role.slug == 'nvl72_poweredgexe9712'], key=natural_rack_key)


def gb300_trays_by_rack(racks: list[Rack]) -> dict[int, list[Device]]:
    rack_ids = [rack.pk for rack in racks]
    grouped = defaultdict(list)
    for device in (
        Device.objects.filter(site__slug=MAD_SITE_SLUG, rack_id__in=rack_ids, device_type__slug='gb300ct')
        .select_related('rack')
        .order_by('rack__name', 'position', 'name')
    ):
        grouped[device.rack_id].append(device)
    return grouped


def be_leaf_devices_by_su() -> dict[int, dict[tuple[int, int], Device]]:
    grouped: dict[int, dict[tuple[int, int], Device]] = defaultdict(dict)
    for device in (
        Device.objects.filter(site__slug=MAD_SITE_SLUG, role__slug='be-leaf-switch')
        .prefetch_related('tags')
        .select_related('rack')
        .order_by('name')
    ):
        su = su_slug(device.tags.all())
        if su is None:
            continue
        source = (device.local_context_data.get('madison_workbook') or {}).get('source_label') or device.description
        match = BE_LEAF_PATTERN.search(source or '')
        if not match:
            continue
        grouped[su][(int(match.group('plane')), int(match.group('switch')))] = device
    return grouped


def shuffle_candidates_by_su_and_group() -> dict[tuple[int, int, str], list[Device]]:
    leaf_devices = Device.objects.filter(site__slug=MAD_SITE_SLUG, role__slug='be-leaf-switch').prefetch_related('tags').select_related('rack')
    rack_su = {}
    for device in leaf_devices:
        su = su_slug(device.tags.all())
        if su is not None and device.rack_id:
            rack_su[device.rack_id] = su

    grouped = defaultdict(list)
    for box in (
        Device.objects.filter(site__slug=MAD_SITE_SLUG, device_type__slug='sb', rack_id__in=rack_su)
        .select_related('rack')
        .order_by('rack__name', 'position', 'name')
    ):
        metadata = box.local_context_data.get(SHUFFLE_CONTAINMENT_MARKER) or {}
        nic = metadata.get('nic_index_zero')
        side = metadata.get('side') or ''
        if nic is None or side not in {'A', 'B'}:
            continue
        grouped[(rack_su[box.rack_id], int(nic), side)].append(box)
    return grouped


def endpoint_for_device_interface(device: Device, interface_name: str, mpo: int, indexes) -> Endpoint | None:
    if (device.pk, interface_name) not in indexes['interfaces']:
        return None
    try:
        return mpo_endpoint_for_device_port(indexes['fabric'], device.name, interface_name, mpo)
    except Endpoint.DoesNotExist:
        return None


def endpoint_for_front_port(device: Device, port_name: str, indexes) -> Endpoint | None:
    if (device.pk, port_name) not in indexes['front_ports']:
        return None
    try:
        return endpoint_for_device_port(indexes['fabric'], device.name, port_name)
    except Endpoint.DoesNotExist:
        return None


def build_indexes():
    fabric = get_fabric()
    interfaces = {
        (interface.device_id, interface.name): interface
        for interface in Interface.objects.filter(device__site__slug=MAD_SITE_SLUG, type__icontains='osfp')
    }
    front_ports = {
        (port.device_id, port.name): port
        for port in FrontPort.objects.filter(
            device__site__slug=MAD_SITE_SLUG,
            device__device_type__slug='shuffle-cassette-2x2-mpo',
        )
    }
    return {
        'fabric': fabric,
        'interfaces': interfaces,
        'front_ports': front_ports,
    }


def main() -> None:
    worksheet_rows = read_worksheet_rows()
    racks = racks_by_su()
    leaf_devices = be_leaf_devices_by_su()
    shuffle_candidates = shuffle_candidates_by_su_and_group()
    indexes = build_indexes()
    tray_devices_by_rack = {}
    for su, su_racks in racks.items():
        gb_racks = gb300_racks_for_su(su_racks)
        tray_devices_by_rack[su] = gb300_trays_by_rack(gb_racks)

    rows = []
    counters = Counter()
    for su in sorted(racks):
        gb_racks = gb300_racks_for_su(racks[su])
        if len(gb_racks) != 7:
            counters['skipped_su_without_7_gb300_racks'] += 1
            continue
        if len(leaf_devices.get(su, {})) != 16:
            counters['skipped_su_without_16_be_leaf_switches'] += 1
            continue

        for pattern in worksheet_rows:
            status = []
            gb_rack = gb_racks[pattern['nvl72_rack_ordinal'] - 1]
            gb_trays = tray_devices_by_rack[su][gb_rack.pk]
            gb_device = gb_trays[pattern['compute_tray_ordinal'] - 1] if len(gb_trays) >= pattern['compute_tray_ordinal'] else None
            gb_endpoint = None
            if gb_device:
                gb_endpoint = endpoint_for_device_interface(gb_device, f"osfp{pattern['cx8']}", pattern['gb300_mpo'], indexes)
            if gb_endpoint is None:
                status.append('missing_gb300_mpo_endpoint')

            leaf_device = leaf_devices[su].get((pattern['plane'], pattern['leaf_switch']))
            leaf_endpoint = None
            if leaf_device:
                leaf_endpoint = endpoint_for_device_interface(leaf_device, f"swp{pattern['leaf_cage']}", pattern['leaf_mpo'], indexes)
            if leaf_endpoint is None:
                status.append('missing_leaf_mpo_endpoint')

            side = 'A' if pattern['plane'] in {1, 2} else 'B'
            nic_index_zero = pattern['leaf_switch'] - 1
            candidates = shuffle_candidates.get((su, nic_index_zero, side), [])
            matching_label_candidates = []
            for candidate in candidates:
                logical = (candidate.local_context_data.get(SHUFFLE_CONTAINMENT_MARKER) or {}).get('logical_shuffle_box')
                if logical == pattern['shuffle_box_ordinal']:
                    matching_label_candidates.append(candidate)
            if not candidates:
                status.append('missing_shuffle_candidate_group')
            elif not matching_label_candidates:
                status.append('shuffle_worksheet_box_number_not_in_elevation_labels')
            elif len(matching_label_candidates) > 1:
                status.append('ambiguous_shuffle_label_candidates')

            counters['rows_total'] += 1
            if not status:
                counters['rows_fully_resolved_by_label'] += 1
            elif gb_endpoint and leaf_endpoint and candidates:
                counters['rows_active_endpoints_resolved_shuffle_candidate_ambiguous'] += 1
            rows.append(
                {
                    'su': su,
                    **pattern,
                    'gb300_rack': gb_rack.name,
                    'gb300_device': gb_device.name if gb_device else '',
                    'gb300_interface': f"osfp{pattern['cx8']}",
                    'gb300_mpo_endpoint_id': gb_endpoint.pk if gb_endpoint else '',
                    'gb300_mpo_endpoint': gb_endpoint.address if gb_endpoint else '',
                    'leaf_device': leaf_device.name if leaf_device else '',
                    'leaf_interface': f"swp{pattern['leaf_cage']}",
                    'leaf_mpo_endpoint_id': leaf_endpoint.pk if leaf_endpoint else '',
                    'leaf_mpo_endpoint': leaf_endpoint.address if leaf_endpoint else '',
                    'shuffle_candidate_side': side,
                    'shuffle_candidate_nic_index_zero': nic_index_zero,
                    'shuffle_candidate_boxes': '|'.join(candidate.name for candidate in candidates),
                    'shuffle_label_matching_candidate_boxes': '|'.join(candidate.name for candidate in matching_label_candidates),
                    'status': ';'.join(status) or 'fully_resolved_by_label',
                }
            )

    path = output_path()
    with path.open('w', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()) if rows else ['status'])
        writer.writeheader()
        writer.writerows(rows)

    print('Madison fiber path resolution report complete.')
    print(f'worksheet_rows_per_su={len(worksheet_rows)}')
    print(f'report_rows={len(rows)}')
    print(f'wrote_csv={path}')
    for key in sorted(counters):
        print(f'{key}={counters[key]}')


main()
