from __future__ import annotations

import argparse
import csv
import json
import os
import re
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


DEFAULT_WORKBOOK_GLOB = 'nscale-nc-18k-fiber-bom-*.xlsx'
DEFAULT_OUTPUT_DIR = Path(__file__).resolve().parents[1] / 'data' / 'generated'


def default_workbook_path() -> Path:
    env_path = os.environ.get('MADISON_FIBER_BOM_WORKBOOK')
    if env_path:
        return Path(env_path)
    source_dir = Path(__file__).resolve().parents[1] / 'data' / 'source'
    matches = sorted(source_dir.glob(DEFAULT_WORKBOOK_GLOB))
    if not matches:
        raise RuntimeError(f'Madison fiber BOM workbook not found with pattern {source_dir / DEFAULT_WORKBOOK_GLOB}')
    if len(matches) > 1:
        raise RuntimeError(f'Multiple Madison fiber BOM workbooks match {source_dir / DEFAULT_WORKBOOK_GLOB}: {matches}')
    return matches[0]


@dataclass(frozen=True)
class BomRow:
    source_sheet: str
    source_row: int
    domain: str
    segment: str
    section: str
    source: str
    dest_rack: str
    cable_type: str
    fiber_count: int
    length_m: float
    bracket_length_m: float | None
    src_leg_m: float | None
    dest_leg_m: float | None
    trunk_qty: int
    active_mpo_per_trunk: int | None
    spare_mpo_per_trunk: int | None
    active_mpo_qty: int
    carried_cable_qty: int
    notes: str


def text(value: Any) -> str:
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def number(value: Any) -> float | None:
    if value in (None, ''):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    match = re.search(r'-?\d+(?:\.\d+)?', str(value))
    return float(match.group(0)) if match else None


def integer(value: Any) -> int:
    parsed = number(value)
    return int(parsed or 0)


def fiber_count(cable_type: str) -> int:
    match = re.search(r'(\d+)\s*f', cable_type.lower())
    return int(match.group(1)) if match else 0


def mpo_profile(fibers: int) -> tuple[int | None, int | None]:
    if fibers == 96:
        return 10, 2
    if fibers == 72:
        return 8, 1
    if fibers == 64:
        return 8, 0
    return None, None


def clean_header(value: Any) -> str:
    return re.sub(r'\s+', ' ', text(value)).strip()


def row_dict(headers: list[str], row: tuple[Any, ...]) -> dict[str, Any]:
    return {header: row[index] if index < len(row) else None for index, header in enumerate(headers)}


def data_rows(ws, header_row: int) -> list[tuple[int, dict[str, Any]]]:
    headers = [clean_header(cell.value) for cell in ws[header_row]]
    rows = []
    for row_index, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        if not any(value is not None for value in row):
            continue
        rows.append((row_index, row_dict(headers, row)))
    return rows


def add_trunk_rows(
    rows: list[BomRow],
    *,
    source_sheet: str,
    source_row: int,
    domain: str,
    segment: str,
    section: str,
    source: str,
    dest_rack: str,
    length_m: float,
    bracket_length_m: float | None,
    src_leg_m: float | None,
    dest_leg_m: float | None,
    trunk_specs: list[tuple[str, int]],
    notes: str = '',
) -> None:
    for cable_type, trunk_qty in trunk_specs:
        if trunk_qty <= 0:
            continue
        fibers = fiber_count(cable_type)
        active_mpo, spare_mpo = mpo_profile(fibers)
        active_qty = trunk_qty * (active_mpo or 0)
        rows.append(
            BomRow(
                source_sheet=source_sheet,
                source_row=source_row,
                domain=domain,
                segment=segment,
                section=section,
                source=source,
                dest_rack=dest_rack,
                cable_type=cable_type,
                fiber_count=fibers,
                length_m=length_m,
                bracket_length_m=bracket_length_m,
                src_leg_m=src_leg_m,
                dest_leg_m=dest_leg_m,
                trunk_qty=trunk_qty,
                active_mpo_per_trunk=active_mpo,
                spare_mpo_per_trunk=spare_mpo,
                active_mpo_qty=active_qty,
                carried_cable_qty=active_qty,
                notes=notes,
            )
        )


def parse_ew_detailed(wb) -> list[BomRow]:
    ws = wb['E-W (CIN) Detailed BOM']
    output: list[BomRow] = []
    for row_index, row in data_rows(ws, 4):
        section = text(row.get('Section'))
        if not section or 'subtotal' in section.lower() or 'grand total' in section.lower() or '\u2192' in section and row.get('Length (m)') is None:
            continue
        length_m = number(row.get('Length (m)'))
        if length_m is None:
            continue
        add_trunk_rows(
            output,
            source_sheet=ws.title,
            source_row=row_index,
            domain='E-W',
            segment='node_to_shuffle',
            section=section,
            source=text(row.get('Source')),
            dest_rack=text(row.get('Dest Rack')),
            length_m=length_m,
            bracket_length_m=None,
            src_leg_m=None,
            dest_leg_m=None,
            trunk_specs=[
                ('96f SM MPO8 trunk', integer(row.get('96f Trunks'))),
                ('72f SM MPO8 trunk', integer(row.get('72f Trunks'))),
            ],
        )
    return output


def parse_ns_detailed(wb) -> list[BomRow]:
    ws = wb['N-S Detailed  BOM']
    output: list[BomRow] = []
    for row_index, row in data_rows(ws, 4):
        section = text(row.get('Section'))
        if not section or 'subtotal' in section.lower() or 'grand total' in section.lower():
            continue
        length_m = number(row.get('Length (m)'))
        if length_m is None:
            continue
        if section.startswith('In-Rack') or section.startswith('Cross-Rack'):
            output.append(
                BomRow(
                    source_sheet=ws.title,
                    source_row=row_index,
                    domain='N-S',
                    segment='leaf_to_spine',
                    section=section,
                    source=text(row.get('Source')),
                    dest_rack=text(row.get('Dest Rack')),
                    cable_type='SM MPO8 patch',
                    fiber_count=8,
                    length_m=length_m,
                    bracket_length_m=None,
                    src_leg_m=None,
                    dest_leg_m=None,
                    trunk_qty=integer(row.get('Total Cables')),
                    active_mpo_per_trunk=1,
                    spare_mpo_per_trunk=0,
                    active_mpo_qty=integer(row.get('Total Cables')),
                    carried_cable_qty=integer(row.get('Total Cables')),
                    notes='Point-to-point MPO8 patch row from N/S detailed BOM.',
                )
            )
            continue
        if 'SU' not in section:
            continue
        add_trunk_rows(
            output,
            source_sheet=ws.title,
            source_row=row_index,
            domain='N-S',
            segment='node_to_leaf',
            section=section,
            source=text(row.get('Source')),
            dest_rack=text(row.get('Dest Rack')),
            length_m=length_m,
            bracket_length_m=None,
            src_leg_m=None,
            dest_leg_m=None,
            trunk_specs=[
                ('96f SM MPO8 trunk', integer(row.get('96f Trunks'))),
                ('64f SM MPO8 trunk', integer(row.get('64f Trunks'))),
            ],
        )
    return output


def parse_ew_patch_rows(wb) -> list[BomRow]:
    ws = wb['E-W (CIN) Consolidated BOM']
    output: list[BomRow] = []
    for row_index, values in enumerate(ws.iter_rows(values_only=True), start=1):
        first = text(values[0] if values else None)
        if first not in {'Shuffle\u2192Leaf', 'Shuffle\u2192Spine'}:
            continue
        source = text(values[1] if len(values) > 1 else '')
        dest = text(values[2] if len(values) > 2 else '')
        cable_type = text(values[3] if len(values) > 3 else '')
        qty = integer(values[4] if len(values) > 4 else 0)
        length = number(cable_type) or 0
        output.append(
            BomRow(
                source_sheet=ws.title,
                source_row=row_index,
                domain='E-W',
                segment='shuffle_to_leaf' if first == 'Shuffle\u2192Leaf' else 'shuffle_to_spine',
                section=first,
                source=source,
                dest_rack=dest,
                cable_type=cable_type or 'SMF MPO8 patch',
                fiber_count=8,
                length_m=length,
                bracket_length_m=None,
                src_leg_m=None,
                dest_leg_m=None,
                trunk_qty=qty,
                active_mpo_per_trunk=1,
                spare_mpo_per_trunk=0,
                active_mpo_qty=qty,
                carried_cable_qty=qty,
                notes='Point-to-point MPO8 patch row from E-W consolidated BOM.',
            )
        )
    return output


def parse_alignment_rows(wb) -> list[BomRow]:
    align_sheets = [name for name in wb.sheetnames if name.startswith('NC 18k Fiber BOM Align')]
    if len(align_sheets) != 1:
        raise RuntimeError(f'Expected exactly one fiber BOM alignment sheet, found: {align_sheets}')
    ws = wb[align_sheets[0]]
    output: list[BomRow] = []
    current_segment = ''
    for row_index, values in enumerate(ws.iter_rows(values_only=True), start=1):
        first = text(values[0] if values else None)
        if first in {'EW BE LEAF TO SHUFFLE (SPINE)', 'N-S SPINE TO CORE'}:
            current_segment = 'be_leaf_to_shuffle_spine' if first.startswith('EW') else 'spine_to_core'
            continue
        if not first or first == 'Cable Type' or first == 'Grand Total' or current_segment == '':
            continue
        qty = integer(values[5] if len(values) > 5 else 0)
        length = number(values[4] if len(values) > 4 else None)
        if qty <= 0 or length is None:
            continue
        fibers = fiber_count(first)
        output.append(
            BomRow(
                source_sheet=ws.title,
                source_row=row_index,
                domain='E-W' if current_segment.startswith('be_') else 'N-S',
                segment=current_segment,
                section=current_segment,
                source='',
                dest_rack='',
                cable_type=first,
                fiber_count=fibers,
                length_m=length,
                bracket_length_m=number(values[1] if len(values) > 1 else None),
                src_leg_m=number(values[2] if len(values) > 2 else None),
                dest_leg_m=number(values[3] if len(values) > 3 else None),
                trunk_qty=qty,
                active_mpo_per_trunk=None,
                spare_mpo_per_trunk=None,
                active_mpo_qty=qty,
                carried_cable_qty=qty,
                notes='Alignment/material row; kept separate from endpoint-ready detailed rows.',
            )
        )
    return output


def parse_bom(path: Path) -> list[BomRow]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows: list[BomRow] = []
    rows.extend(parse_ew_detailed(workbook))
    rows.extend(parse_ew_patch_rows(workbook))
    rows.extend(parse_ns_detailed(workbook))
    rows.extend(parse_alignment_rows(workbook))
    return rows


def write_outputs(rows: list[BomRow], output_dir: Path) -> tuple[Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    json_path = output_dir / 'madison_fiber_bom_manifest.json'
    csv_path = output_dir / 'madison_fiber_bom_manifest.csv'
    json_path.write_text(json.dumps([asdict(row) for row in rows], indent=2) + '\n')

    with csv_path.open('w', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=list(asdict(rows[0]).keys()))
        writer.writeheader()
        for row in rows:
            writer.writerow(asdict(row))

    return json_path, csv_path


def print_report(rows: list[BomRow]) -> None:
    endpoint_rows = [row for row in rows if 'Alignment/material row' not in row.notes]
    by_segment = defaultdict(lambda: Counter({'trunks': 0, 'active_mpo': 0, 'carried': 0}))
    by_type = Counter()
    for row in endpoint_rows:
        by_segment[row.segment]['trunks'] += row.trunk_qty
        by_segment[row.segment]['active_mpo'] += row.active_mpo_qty
        by_segment[row.segment]['carried'] += row.carried_cable_qty
        by_type[row.cable_type] += row.trunk_qty

    print('Madison fiber BOM manifest')
    print(f'total_rows={len(rows)}')
    print(f'endpoint_or_patch_rows={len(endpoint_rows)}')
    print('\nEndpoint/patch totals by segment:')
    for segment in sorted(by_segment):
        values = by_segment[segment]
        print(f'  {segment}: trunk_or_patch_qty={values["trunks"]}, active_mpo_or_patch_qty={values["active_mpo"]}')

    print('\nEndpoint/patch totals by cable type:')
    for cable_type, qty in sorted(by_type.items()):
        print(f'  {cable_type}: {qty}')

    print('\nLength distribution by segment:')
    length_counter = defaultdict(Counter)
    for row in endpoint_rows:
        length_counter[row.segment][row.length_m] += row.trunk_qty
    for segment in sorted(length_counter):
        pretty = ', '.join(f'{length:g}m={qty}' for length, qty in sorted(length_counter[segment].items()))
        print(f'  {segment}: {pretty}')


def main() -> None:
    parser = argparse.ArgumentParser(description='Normalize the Madison NC fiber BOM workbook into a generated CSV/JSON manifest.')
    parser.add_argument('--workbook', type=Path, default=None)
    parser.add_argument('--output-dir', type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument('--no-write', action='store_true')
    args = parser.parse_args()

    rows = parse_bom(args.workbook or default_workbook_path())
    print_report(rows)
    if not args.no_write:
        json_path, csv_path = write_outputs(rows, args.output_dir)
        print(f'\nwrote_json={json_path}')
        print(f'wrote_csv={csv_path}')


if __name__ == '__main__':
    main()
