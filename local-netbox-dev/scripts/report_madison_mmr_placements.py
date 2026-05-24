#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import json
import os
import re
from collections import Counter
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


DEFAULT_WORKBOOK = Path(
    os.environ.get(
        'MADISON_LAYOUT_WORKBOOK',
        Path(__file__).resolve().parents[1] / 'data' / 'source' / 'madison-layout-workbook.xlsx',
    )
)
DEFAULT_OUTPUT_DIR = Path(__file__).resolve().parents[1] / 'data' / 'generated'
MMR_SHEET = 'MMRs'

KNOWN_DEVICE_TYPE_SLUGS = {
    'generic-proxmox-node',
    'nokia-breakout-40x10',
    'nokia-ixr-d5',
    'nokia-ixs-a1',
    'nokia-sr1',
    'opengear-om2224-24e-l',
    'palo-alto-pa-1420',
    'palo-alto-pa-550',
}

MMR_COLUMNS = (
    ('mmr1', 'MMR1', 2, 3),
    ('mmr2', 'MMR2', 6, 7),
)


@dataclass(frozen=True)
class MMRPlacement:
    mmr_slug: str
    mmr_name: str
    source_sheet: str
    source_cell: str
    ru_top: int
    ru_bottom: int
    height_u: int
    source_label: str
    normalized_label: str
    device_type_slug: str
    device_role_slug: str
    device_type_defined: bool
    notes: str


def clean_label(value: Any) -> str:
    return ' '.join(str(value or '').replace('\n', ' ').split())


def int_ru(value: Any) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and 1 <= value <= 48:
        return int(value)
    return None


def classify_label(label: str) -> tuple[str, str, str, list[str]]:
    lower = label.lower()
    notes: list[str] = ['MMR sheet placement; not yet consumed by Madison device-instance seeders.']

    if 'fiber panel' in lower:
        notes.append('Fiber panel label is unresolved; final panel/trunk mapping is deferred.')
        return 'unresolved-fiber-panel', 'fiber-panel', 'Fiber Panel', notes
    if 'palo alto 1420' in lower:
        return 'palo-alto-pa-1420', 'nscale-firewall', 'Palo Alto PA-1420', notes
    if 'palo alto 550' in lower:
        return 'palo-alto-pa-550', 'nscale-firewall', 'Palo Alto PA-550', notes
    if 'opengear' in lower or 'om2224' in lower:
        return 'opengear-om2224-24e-l', 'console-server', 'Opengear OM2224', notes
    if 'nokia sr1' in lower or 'nokia sr-1' in lower:
        return 'nokia-sr1', 'edge-switch', 'Nokia SR1', notes
    if 'nokia ixr' in lower:
        return 'nokia-ixr-d5', 'edge-switch', 'Nokia IXR-D5', notes
    if 'nokia breakout' in lower:
        return 'nokia-breakout-40x10', 'passive-breakout', 'Nokia Breakout 40G to 4x10G', notes
    if 'ixs-a1' in lower:
        return 'nokia-ixs-a1', 'oob-leaf', 'Nokia IXS-A1', notes
    if re.search(r'\bproxmox\b', lower):
        return 'generic-proxmox-node', 'management-server', 'Proxmox Management Node', notes

    return '', '', lower, [*notes, 'No classifier matched this MMR workbook label.']


def parse_mmr_placements(workbook_path: Path) -> list[MMRPlacement]:
    workbook = load_workbook(workbook_path, data_only=True, read_only=False)
    if MMR_SHEET not in workbook.sheetnames:
        return []
    worksheet = workbook[MMR_SHEET]
    placements: list[MMRPlacement] = []

    for mmr_slug, mmr_name, ru_column, label_column in MMR_COLUMNS:
        for row_number in range(1, worksheet.max_row + 1):
            ru = int_ru(worksheet.cell(row_number, ru_column).value)
            label_cell = worksheet.cell(row_number, label_column)
            source_label = clean_label(label_cell.value)
            if ru is None or not source_label:
                continue
            device_type_slug, device_role_slug, normalized_label, notes = classify_label(source_label)
            placements.append(
                MMRPlacement(
                    mmr_slug=mmr_slug,
                    mmr_name=mmr_name,
                    source_sheet=worksheet.title,
                    source_cell=label_cell.coordinate,
                    ru_top=ru,
                    ru_bottom=ru,
                    height_u=1,
                    source_label=source_label,
                    normalized_label=normalized_label,
                    device_type_slug=device_type_slug,
                    device_role_slug=device_role_slug,
                    device_type_defined=bool(device_type_slug and device_type_slug in KNOWN_DEVICE_TYPE_SLUGS),
                    notes=' '.join(notes),
                )
            )

    placements.sort(key=lambda item: (item.mmr_slug, -item.ru_top, item.source_cell))
    return placements


def write_outputs(placements: list[MMRPlacement], output_dir: Path) -> tuple[Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    json_path = output_dir / 'madison_mmr_placement_manifest.json'
    csv_path = output_dir / 'madison_mmr_placement_manifest.csv'

    json_path.write_text(json.dumps([asdict(placement) for placement in placements], indent=2) + '\n')
    fieldnames = list(asdict(placements[0]).keys()) if placements else [field.name for field in MMRPlacement.__dataclass_fields__.values()]
    with csv_path.open('w', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for placement in placements:
            writer.writerow(asdict(placement))

    return json_path, csv_path


def print_report(placements: list[MMRPlacement], *, include_samples: bool) -> None:
    unmatched = [placement for placement in placements if not placement.device_type_slug]
    missing_definitions = sorted(
        {
            placement.device_type_slug
            for placement in placements
            if placement.device_type_slug and not placement.device_type_defined
        }
    )

    print('Madison MMR placement manifest')
    print(f'mmr_placement_rows={len(placements)}')
    print(f'mmr_racks={len({placement.mmr_slug for placement in placements})}')
    print(f'unmatched_label_placements={len(unmatched)}')

    print('\nBy MMR:')
    for mmr_slug, count in sorted(Counter(placement.mmr_slug for placement in placements).items()):
        print(f'  {mmr_slug}: {count}')

    print('\nBy device type:')
    for slug, count in sorted(Counter(placement.device_type_slug or 'unmatched' for placement in placements).items()):
        print(f'  {slug}: {count}')

    print('\nBy device role:')
    for slug, count in sorted(Counter(placement.device_role_slug or 'unmatched' for placement in placements).items()):
        print(f'  {slug}: {count}')

    if missing_definitions:
        print('\nDevice types still missing from staged definitions:')
        for slug in missing_definitions:
            count = sum(1 for placement in placements if placement.device_type_slug == slug)
            print(f'  {slug}: {count}')

    if unmatched:
        print('\nUnmatched MMR labels:')
        for placement in unmatched:
            print(f'  {placement.mmr_name} RU{placement.ru_top}: {placement.source_label} ({placement.source_sheet}!{placement.source_cell})')

    if include_samples:
        print('\nMMR samples:')
        for placement in placements:
            print(
                f'  {placement.mmr_name} RU{placement.ru_top}: '
                f'{placement.source_label} -> {placement.device_type_slug or "unmatched"}/{placement.device_role_slug or "unmatched"}'
            )


def main() -> None:
    parser = argparse.ArgumentParser(description='Extract read-only Madison MMR placements from the Google layout workbook.')
    parser.add_argument('--workbook', type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument('--output-dir', type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument('--include-samples', action='store_true')
    parser.add_argument('--no-write', action='store_true')
    args = parser.parse_args()

    placements = parse_mmr_placements(args.workbook)
    print_report(placements, include_samples=args.include_samples)

    if not args.no_write:
        json_path, csv_path = write_outputs(placements, args.output_dir)
        print(f'\nwrote_json={json_path}')
        print(f'wrote_csv={csv_path}')


if __name__ == '__main__':
    main()
