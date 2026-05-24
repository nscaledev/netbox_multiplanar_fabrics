#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import json
import os
import re
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from pathlib import Path

from openpyxl import load_workbook


DEFAULT_WORKBOOK = Path(
    os.environ.get(
        'MADISON_LAYOUT_WORKBOOK',
        Path(__file__).resolve().parents[1] / 'data' / 'source' / 'madison-layout-workbook.xlsx',
    )
)
DEFAULT_OUTPUT_DIR = Path(__file__).resolve().parents[1] / 'data' / 'generated'
NC_SU_MAPPING_SHEET = 'NC SU Mapping'

OCCUPIED_SLOT_ROWS = (4, 8, 12, 16)
SLOT_PATTERN = re.compile(r'^[A-X]\d{1,2}$')
SU_PATTERN = re.compile(r'\bSU\s*(\d+)(?:\s*/\s*(\d+))?\b', re.IGNORECASE)
PLANE_PATTERN = re.compile(r'\bP\s*(\d+)\b', re.IGNORECASE)

ROLE_SLUGS = {
    'nvl72_poweredgexe9712': 'NVL72_PowerEdgeXE9712',
    'madison-t1-ew-plane-1-2': 'GS001 T1 E-W Plane 1-2',
    'madison-t1-ew-plane-3-4': 'GS001 T1 E-W Plane 3-4',
    'madison-t1-t2-ns': 'GS001 T1/T2 N-S',
    'madison-be-spine-plane-1': 'GS001 BE Spine Plane 1',
    'madison-be-spine-plane-2': 'GS001 BE Spine Plane 2',
    'madison-be-spine-plane-3': 'GS001 BE Spine Plane 3',
    'madison-be-spine-plane-4': 'GS001 BE Spine Plane 4',
    'madison-fe-core-edge': 'GS001 FE Core + Edge',
    'madison-control': 'GS001 Control',
    'madison-t1-t2-storage': 'GS001 T1/T2 Storage',
    'empty-or-spacer': 'Empty or Spacer',
    'unknown': 'Unknown',
}


@dataclass(frozen=True)
class RackSlot:
    physical_slot: str
    workbook_cell: str
    location: str
    source_label: str
    role_slug: str
    role_name: str
    scalable_units: list[int]
    backend_plane: int | None
    rack_level_su_tags: list[str]
    dimensions: str | None
    notes: str


def clean_label(value) -> str:
    return ' '.join(str(value or '').replace('\n', ' ').split())


def location_for_slot(physical_slot: str) -> str:
    row_letter = physical_slot[0]
    return 'Data Hall 2' if 'A' <= row_letter <= 'L' else 'Data Hall 1'


def classify_role(physical_slot: str, label: str) -> str:
    normalized = label.lower()
    slot_number = int(physical_slot[1:])

    if not normalized:
        return 'empty-or-spacer'
    if 'gpu' in normalized and 'rack' in normalized:
        return 'nvl72_poweredgexe9712'
    if 'be - leaf' in normalized:
        return 'madison-t1-ew-plane-1-2' if slot_number % 2 == 1 else 'madison-t1-ew-plane-3-4'
    if 'fe - leaf' in normalized:
        return 'madison-t1-t2-ns'
    if 'be - spine' in normalized:
        plane_match = PLANE_PATTERN.search(label)
        if plane_match:
            return f'madison-be-spine-plane-{plane_match.group(1)}'
        return 'unknown'
    if 'fe core' in normalized:
        return 'madison-fe-core-edge'
    if 'control plane' in normalized:
        return 'madison-control'
    if 'bm storage' in normalized:
        return 'madison-t1-t2-storage'
    return 'unknown'


def rack_su_tags_for(role_slug: str, scalable_units: list[int]) -> list[str]:
    if not scalable_units:
        return []
    if role_slug not in {
        'nvl72_poweredgexe9712',
        'madison-t1-ew-plane-1-2',
        'madison-t1-ew-plane-3-4',
        'madison-t1-t2-ns',
    }:
        return []
    return [f'nv_su_{scalable_unit}' for scalable_unit in scalable_units]


def parse_scalable_units(label: str) -> list[int]:
    scalable_units = set()
    for first_su, second_su in SU_PATTERN.findall(label):
        scalable_units.add(int(first_su))
        if second_su:
            scalable_units.add(int(second_su))
    return sorted(scalable_units)


def parse_manifest(workbook_path: Path) -> list[RackSlot]:
    workbook = load_workbook(workbook_path, data_only=True)
    worksheet = workbook[NC_SU_MAPPING_SHEET]
    slots: list[RackSlot] = []

    for row_number in OCCUPIED_SLOT_ROWS:
        for cell in worksheet[row_number]:
            physical_slot = cell.value
            if not (isinstance(physical_slot, str) and SLOT_PATTERN.fullmatch(physical_slot.strip())):
                continue

            source_label = clean_label(worksheet.cell(cell.row - 1, cell.column).value)
            dimensions = worksheet.cell(cell.row - 2, cell.column).value
            role_slug = classify_role(physical_slot.strip(), source_label)
            scalable_units = parse_scalable_units(source_label)
            plane_match = PLANE_PATTERN.search(source_label)
            backend_plane = int(plane_match.group(1)) if plane_match and 'be - spine' in source_label.lower() else None

            notes = []
            if role_slug == 'empty-or-spacer':
                notes.append('Physical slot is present in workbook but is empty/spacer.')
            if role_slug == 'unknown':
                notes.append('Workbook label did not match a known rack-role classifier.')
            if role_slug == 'madison-t1-t2-ns' and len(scalable_units) > 1:
                notes.append('Shared FE leaf rack; apply all listed SU tags because all modeled resources are SU-scoped.')
            if role_slug.startswith('madison-be-spine-plane-'):
                notes.append('Backend plane is not the same construct as SU.')

            slots.append(
                RackSlot(
                    physical_slot=physical_slot.strip(),
                    workbook_cell=cell.coordinate,
                    location=location_for_slot(physical_slot.strip()),
                    source_label=source_label,
                    role_slug=role_slug,
                    role_name=ROLE_SLUGS.get(role_slug, ROLE_SLUGS['unknown']),
                    scalable_units=scalable_units,
                    backend_plane=backend_plane,
                    rack_level_su_tags=rack_su_tags_for(role_slug, scalable_units),
                    dimensions=str(dimensions) if dimensions is not None else None,
                    notes=' '.join(notes),
                )
            )

    return sorted(slots, key=lambda slot: (slot.physical_slot[0], int(slot.physical_slot[1:])))


def write_outputs(slots: list[RackSlot], output_dir: Path) -> tuple[Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    json_path = output_dir / 'madison_workbook_rack_manifest.json'
    csv_path = output_dir / 'madison_workbook_rack_manifest.csv'

    json_path.write_text(json.dumps([asdict(slot) for slot in slots], indent=2) + '\n')

    with csv_path.open('w', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=list(asdict(slots[0]).keys()))
        writer.writeheader()
        for slot in slots:
            row = asdict(slot)
            row['scalable_units'] = ','.join(str(su) for su in slot.scalable_units)
            row['rack_level_su_tags'] = ','.join(slot.rack_level_su_tags)
            writer.writerow(row)

    return json_path, csv_path


def print_report(slots: list[RackSlot], *, include_empty: bool) -> None:
    occupied_slots = [slot for slot in slots if slot.role_slug != 'empty-or-spacer']
    tagged_slots = [slot for slot in occupied_slots if slot.rack_level_su_tags]
    shared_su_slots = [slot for slot in occupied_slots if len(slot.scalable_units) > 1]
    unknown_slots = [slot for slot in occupied_slots if slot.role_slug == 'unknown']

    print('Madison workbook rack manifest')
    print(f'total_physical_slots={len(slots)}')
    print(f'occupied_rack_slots={len(occupied_slots)}')
    print(f'empty_or_spacer_slots={len(slots) - len(occupied_slots)}')
    print(f'rack_level_su_tagged_slots={len(tagged_slots)}')
    print(f'rack_level_su_tag_assignments={sum(len(slot.rack_level_su_tags) for slot in tagged_slots)}')
    print(f'shared_su_slots={len(shared_su_slots)}')
    print(f'unknown_role_slots={len(unknown_slots)}')

    print('\nBy location:')
    for location, count in sorted(Counter(slot.location for slot in occupied_slots).items()):
        print(f'  {location}: {count}')

    print('\nBy role:')
    for role_slug, count in sorted(Counter(slot.role_slug for slot in occupied_slots).items()):
        print(f'  {role_slug}: {count}')

    print('\nBy SU tag assignment:')
    su_counts = Counter()
    for slot in tagged_slots:
        for scalable_unit in slot.scalable_units:
            su_counts[scalable_unit] += 1
    for su in sorted(su_counts):
        print(f'  SU{su}: {su_counts[su]}')

    print('\nShared multi-SU tagged slots:')
    for slot in shared_su_slots:
        print(f'  {slot.physical_slot}: {slot.role_slug} {slot.source_label}')

    if unknown_slots:
        print('\nUnknown role slots:')
        for slot in unknown_slots:
            print(f'  {slot.physical_slot}: {slot.source_label}')

    if include_empty:
        print('\nEmpty/spacer slots:')
        for slot in slots:
            if slot.role_slug == 'empty-or-spacer':
                print(f'  {slot.physical_slot}: {slot.workbook_cell}')

    per_row = defaultdict(Counter)
    for slot in slots:
        per_row[slot.physical_slot[0]][slot.role_slug] += 1

    print('\nRow summary:')
    for row_letter in sorted(per_row):
        parts = ', '.join(f'{role}={count}' for role, count in sorted(per_row[row_letter].items()))
        print(f'  Row {row_letter}: {parts}')


def main() -> None:
    parser = argparse.ArgumentParser(description='Build a rack-slot manifest from the Madison workbook NC SU Mapping sheet.')
    parser.add_argument('--workbook', type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument('--output-dir', type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument('--include-empty', action='store_true')
    parser.add_argument('--no-write', action='store_true')
    args = parser.parse_args()

    slots = parse_manifest(args.workbook)
    print_report(slots, include_empty=args.include_empty)

    if not args.no_write:
        json_path, csv_path = write_outputs(slots, args.output_dir)
        print(f'\nwrote_json={json_path}')
        print(f'wrote_csv={csv_path}')


if __name__ == '__main__':
    main()
